import os

from docx import Document

from constant import path_storage


def find_ps(sender_id=int):
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    file_list = os.listdir(path)
    des_file = str(sender_id) + '_S.docx'
    if des_file not in file_list:
        return False
    else:
        return True


def search_str(origin_list, search):
    if isinstance(origin_list, list):
        return search in origin_list


def update_buff(sender_id=int, op_type=str, op_buff=str):
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_S.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_S.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]

    buff_content = table.cell(4, 3).text
    if len(buff_content) > 0 and buff_content[0] == '\n':
        buff_content = buff_content[1:]
    if len(buff_content) > 0 and buff_content[-1] == '\n':
        buff_content = buff_content[:-1]
    buff_list = str(buff_content).split('\n')
    result_buff = ''

    if op_type == '+':
        # 加入新buff
        if search_str(buff_list, op_buff):
            # 已经有该buff
            return 'Have'
        else:
            for each in buff_list:
                result_buff += each + '\n'
            result_buff += str(op_buff)
    elif op_type == '-':
        # 去除已有buff
        if not search_str(buff_list, op_buff):
            # 没有该buff
            return 'Not have'
        else:
            for each in buff_list:
                if each != op_buff:
                    result_buff += each + '\n'
    if len(result_buff) > 0 and result_buff[0] == '\n':
        result_buff = result_buff[1:]
    if len(result_buff) > 0 and result_buff[-1] == '\n':
        result_buff = result_buff[:-1]
    # 将修改后的字符串内容重新写入
    table.cell(4, 3).text = result_buff
    document.save(path)
    return 'Succeed'
