# -*- coding: UTF-8 -*-

# Filename : test.py
# author by : www.runoob.com
import os

from docx import Document

from constant import path_storage


def is_number(s):
    try:
        float(s)
        return True
    except ValueError:
        pass

    try:
        import unicodedata
        unicodedata.numeric(s)
        return True
    except (TypeError, ValueError):
        pass

    return False


def find_ps(sender_id=int):
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    file_list = os.listdir(path)
    des_file = str(sender_id) + '_S.docx'
    if des_file not in file_list:
        return False
    else:
        return True


def change_hp(sender_id, change_health):
    difference = float(change_health)
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_S.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_S.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]
    cur_hp_cell = table.cell(0, 3)
    max_hp_cell = table.cell(0, 10)
    cur_hp = float(cur_hp_cell.text)
    max_hp = float(max_hp_cell.text)

    result_hp = cur_hp + difference
    if result_hp > max_hp:
        result_hp = max_hp  # 不能超过上限
    if result_hp < 0:
        result_hp = 0  # 不能低于0
    cur_hp_cell.text = str(result_hp)
    document.save(path)
    final_hp = table.cell(0, 3).text
    hp_str = final_hp + '/' + str(max_hp)
    return hp_str
