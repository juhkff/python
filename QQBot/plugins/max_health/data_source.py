import os

from docx import Document

from constant import path_storage


def change_max_hp(sender_id, change_health):
    difference = float(change_health)
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_S.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_S.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]
    cur_hp_cell = table.cell(0, 3)
    max_hp_cell = table.cell(0, 10)
    cur_hp = cur_hp_cell.text
    max_hp = float(max_hp_cell.text)

    result_hp = max_hp + difference
    if result_hp < 0:
        result_hp = 0  # 不能低于0
    max_hp_cell.text = str(result_hp)
    document.save(path)
    final_hp = table.cell(0, 10).text
    hp_str = cur_hp + '/' + final_hp
    return hp_str
