import configparser
import os

from docx import Document

from constant import path_settings, path_storage


# 在ini文件中寻找名称
def search_name_in_settings(name):
    con = configparser.ConfigParser()
    con.read(path_settings, encoding='gbk')
    items = con.items('playerProp2')  # 返回元组
    a_dict = dict(items)
    items_real = con.items('playerProp')
    b_dict = dict(items_real)
    # name_list = list(a_dict.values())
    for index in a_dict:
        if a_dict[index] == name:
            return [True, index, str(b_dict[index])]  # 返回True以及该名称对应的key以及实际名字
    for index2 in b_dict:
        if b_dict[index2] == name:
            return [True, index2, name]
    return [False, None, None]


def change_pro(sender_id, change_number, pro_name, pro_index):
    pro_index = int(pro_index)
    difference = int(change_number)
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_C.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_C.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]

    index = 1
    between_index = 0
    between = [4, 3, 4, 1]
    origin_line = 4
    for i in range(1, pro_index):
        index += between[between_index]
        between_index = (between_index + 1) % 4
    table_cell = table.cell(origin_line + int((index - 1) / 12), ((index - 1) % 12) + 1)
    origin_content = table_cell.text

    # origin_content = 原框中的内容
    split_list = str(origin_content).split('(')
    origin_number = int(split_list[0])
    origin_extra = ''
    if len(split_list) > 1:
        origin_extra = split_list[1]  # 注意补(
    result_number = origin_number + difference
    result_content = str(result_number)
    if len(split_list) > 1:
        result_content = result_content + '(' + origin_extra

    table_cell.text = str(result_content)
    document.save(path)
    final_content = table.cell(origin_line + int((index - 1) / 12), ((index - 1) % 12) + 1).text
    return final_content
