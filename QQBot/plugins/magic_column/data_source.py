import os

from docx import Document

from constant import path_storage


def change_magic_column(sender_id, change_column, change_number):
    change_index = None
    difference = int(change_number)
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_S.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_S.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]

    # 不同情况
    if change_column == '1':
        change_index = 0
    elif change_column == '2':
        change_index = 1
    elif change_column == '3':
        change_index = 2
    elif change_column == '4':
        change_index = 4
    elif change_column == '5':
        change_index = 6
    elif change_column == '6':
        change_index = 7
    elif change_column == '7':
        change_index = 8
    elif change_column == '8':
        change_index = 10
    elif change_column == '9':
        change_index = 11

    cur_magic_consume = int(table.cell(3, change_index).text)
    result_int = cur_magic_consume + difference
    nothing = False
    if result_int < 0:  # 低于0无法完成该操作，返回旧值
        nothing = True
        return [cur_magic_consume, nothing]
    table.cell(3, change_index).text = str(result_int)
    document.save(path)
    return [result_int, nothing]
