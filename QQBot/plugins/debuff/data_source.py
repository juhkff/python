import os

from docx import Document

from constant import path_storage


def find_ps(sender_id=int):
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    file_list = os.listdir(path)
    des_file = str(sender_id) + '_S.docx'
    if des_file not in file_list:
        return False
    else:
        return True


def search_str(origin_list, search):
    if isinstance(origin_list, list):
        return search in origin_list


def update_debuff(sender_id=int, op_type=str, op_buff=str):
    path = os.path.join(path_storage, str(sender_id))  # storage/qq/
    des_file = str(sender_id) + '_S.docx'
    path = os.path.join(path, des_file)  # storage/qq/qq_S.docx
    document = Document(path)
    tables = document.tables  # 获取文件中的表格集
    table = tables[0]

    debuff_content = table.cell(5, 3).text
    if len(debuff_content) > 0 and debuff_content[0] == '\n':
        debuff_content = debuff_content[1:]
    if len(debuff_content) > 0 and debuff_content[-1] == '\n':
        debuff_content = debuff_content[:-1]
    debuff_list = str(debuff_content).split('\n')
    result_debuff = ''

    if op_type == '+':
        # 加入新debuff
        if search_str(debuff_list, op_buff):
            # 已经有该debuff
            return 'Have'
        else:
            for each in debuff_list:
                result_debuff += each + '\n'
            result_debuff += str(op_buff)
    elif op_type == '-':
        # 去除已有debuff
        if not search_str(debuff_list, op_buff):
            # 没有该debuff
            return 'Not have'
        else:
            for each in debuff_list:
                if each != op_buff:
                    result_debuff += each + '\n'
    if len(result_debuff) > 0 and result_debuff[0] == '\n':
        result_debuff = result_debuff[1:]
    if len(result_debuff) > 0 and result_debuff[-1] == '\n':
        result_debuff = result_debuff[:-1]
    # 将修改后的字符串内容重新写入
    table.cell(5, 3).text = result_debuff
    document.save(path)
    return 'Succeed'
